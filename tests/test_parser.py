import io
import pytest
import fitz
import docx
from app.services.document_parser import (
    extract_text,
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_txt
)
from app.core.exceptions import (
    DocumentParsingError,
    UnsupportedFileTypeError,
    EmptyDocumentError,
    FileTooLargeError
)
from app.config import settings

@pytest.fixture
def valid_pdf_bytes():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Python FastAPI developer resume details.")
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes

@pytest.fixture
def valid_docx_bytes():
    doc = docx.Document()
    doc.add_paragraph("Experienced Machine Learning engineer with Python experience.")
    # Add a table to test table extraction
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Skills'
    hdr_cells[1].text = 'Docker, AWS'
    
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()

@pytest.fixture
def valid_txt_bytes():
    return "Software Engineer specializing in PostgreSQL, Docker, and Python.".encode("utf-8")

def test_pdf_extraction(valid_pdf_bytes):
    text = extract_text(valid_pdf_bytes, "resume.pdf")
    assert "Python FastAPI developer" in text

def test_docx_extraction(valid_docx_bytes):
    text = extract_text(valid_docx_bytes, "resume.docx")
    assert "Experienced Machine Learning" in text
    assert "Docker, AWS" in text

def test_txt_extraction(valid_txt_bytes):
    text = extract_text(valid_txt_bytes, "resume.txt")
    assert "PostgreSQL, Docker" in text

def test_empty_document():
    with pytest.raises(EmptyDocumentError):
        extract_text(b"", "resume.txt")

def test_empty_pdf_no_text():
    # PDF with page but no text content
    doc = fitz.open()
    doc.new_page()
    pdf_bytes = doc.write()
    doc.close()
    with pytest.raises(EmptyDocumentError):
        extract_text(pdf_bytes, "empty.pdf")

def test_unsupported_extension():
    with pytest.raises(UnsupportedFileTypeError):
        extract_text(b"some bytes", "resume.png")
    with pytest.raises(UnsupportedFileTypeError):
        # Test old doc format exception
        extract_text(b"some bytes", "resume.doc")

def test_corrupted_pdf():
    with pytest.raises(DocumentParsingError):
        extract_text(b"invalid pdf content", "resume.pdf")

def test_corrupted_docx():
    with pytest.raises(DocumentParsingError):
        extract_text(b"invalid docx zip content", "resume.docx")

def test_extremely_large_file(valid_txt_bytes, monkeypatch):
    # Set max file size to 0MB to trigger size validation
    monkeypatch.setattr(settings, "MAX_FILE_SIZE_MB", 0)
    with pytest.raises(FileTooLargeError) as excinfo:
        extract_text(valid_txt_bytes, "resume.txt")
    assert "size exceeds" in str(excinfo.value)
